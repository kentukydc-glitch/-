import json,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table,TableStyleInfo
d=json.load(open('merged.json'));rows=d['rows']
F=Font(name='Arial',size=10);B=Font(name='Arial',size=10,bold=True);H=Font(name='Arial',size=10,bold=True,color='FFFFFF')
HF=PatternFill('solid',fgColor='1F3864');TOP=PatternFill('solid',fgColor='FFF2CC');BLUE=Font(name='Arial',size=10,color='0000FF');GREEN=Font(name='Arial',size=10,color='008000')
WR=Alignment(wrap_text=True,vertical='top')
def s(v):
  if isinstance(v,list): return '; '.join(map(str,v))
  return '' if v is None else str(v)
def moldova(c):
  m=c['moldova_offered'].lower()
  return 'Да' if m.startswith('yes') else ('Не найдено (неясно)' if 'unclear' in m else 'Нет')
RISK={'GSL':'Телефон и присутствие в Москве: до контакта проверить санкционные и репутационные риски',
 'Eurofast':'Последние датированные новости 2022 г.: проверить, активен ли сайт',
 'Vistra':'Поглотила Biz Latin Hub и Tricor; крупная корпорация, выйти на нужного человека сложно',
 'Healy':'Входит в Hawksford; публичной партнёрской программы нет',
 'Virtuzone':'В задании был указан домен vzone.ae, это посторонний сайт; правильный домен virtuzone.com',
 'InCorp':'Не путать с InCorp Services (США, incorp.com), у которой есть реферальная программа 20%',
 'Uniwide':'Общий email не найден; связь через форму или телефоны'}
wb=Workbook()
# --- Методика
m=wb.active;m.title='Методика'
W=[('Релевантность клиентской базы','client_relevance',0.30),('Открытость к партнёрству','partner_openness',0.25),('Возможность по Молдове','moldova_opportunity',0.25),('Доступность контактов','reachability',0.10),('Масштаб','scale',0.10)]
m['A1']='Рейтинг перспективности: веса критериев (синие ячейки можно менять, рейтинг пересчитается)';m['A1'].font=B
m['A2']='Критерий';m['B2']='Вес';m['C2']='Шкала';[setattr(m[c],'font',B) for c in('A2','B2','C2')]
for i,(n,k,w) in enumerate(W,3):
  m[f'A{i}']=n;m[f'B{i}']=w;m[f'B{i}'].font=BLUE;m[f'B{i}'].number_format='0%';m[f'C{i}']='1–5'
m['A8']='Сумма весов';m['B8']='=SUM(B3:B7)';m['B8'].number_format='0%'
notes=['','Как читать «Возможность по Молдове»: 5 = Молдовы нет, но покрыты соседние рынки (Румыния, Украина, CEE) и компания работает через локальных партнёров; 3–4 = Молдова уже предлагается и, вероятно, нужен местный юрист.',
'Группа A: в опубликованном каталоге Молдова отсутствует или не найдена. Группа B: Молдова уже предлагается.',
'ОГРАНИЧЕНИЕ: сетевая политика среды исследования блокировала прямой доступ к сайтам компаний (HTTP 403 / DNS). Все факты взяты из поисковой выдачи по официальным доменам, URL источников указаны. Перед контактом откройте ссылки и сверьте данные. Контакты из сторонних каталогов помечены.',
'Ничего не выдумано: если факт не подтверждён, стоит «not found» / «not verified». Письма НЕ отправлялись.',
'Дата исследования: 2026-10-08.']
for i,t in enumerate(notes,9): m[f'A{i}']=t;m[f'A{i}'].alignment=WR;m.merge_cells(f'A{i}:F{i}');m.row_dimensions[i].height=45 if t else 10
m.column_dimensions['A'].width=40;m.column_dimensions['B'].width=10
for r in m.iter_rows():
  for c in r:
    if c.font!=B and c.font!=BLUE: c.font=F
for i in range(3,9):m[f'B{i}'].font=BLUE if i<8 else F
# --- Рейтинг
r=wb.create_sheet('Рейтинг')
cols=[('Ранг',6),('TOP-5',7),('Компания',26),('Штаб-квартира',16),('Сайт',24),('Сегмент',20),('Группа',10),('Молдова в каталоге',12),('Доказательство по Молдове',45),('Источник (Молдова)',35),
('Юрисдикции (кол-во)',22),('Юрисдикции (примеры)',35),('Источник (юрисдикции)',35),('Партнёрская программа',14),('Детали программы',45),('Источник (программа)',35),
('Email',26),('Телефон',22),('Страница контактов',30),('Адрес',30),('Источник контактов',35),('Почему подходит Legitimus',50),('Риски / проверить',35),('Признаки активности',30),('Уровень проверки',22),
('Клиенты',8),('Партнёрство',9),('Молдова',8),('Доступность',9),('Масштаб',8),('Итоговый балл',9)]
for j,(h,w) in enumerate(cols,1):
  c=r.cell(1,j,h);c.font=H;c.fill=HF;c.alignment=Alignment(wrap_text=True,vertical='center');r.column_dimensions[L(j)].width=w
n=len(rows);last=n+1
for i,c in enumerate(rows,2):
  pp=c['partner_program'].lower()
  vals=[f'=RANK(AE{i},$AE$2:$AE${last})+COUNTIF($AE$2:AE{i},AE{i})-1',f'=IF(A{i}<=5,"TOP-5","")',c['name'],c['hq_country'],c['website'],c['segment'],
   '=IF(H%d="Да","B","A")'%i,moldova(c),s(c['moldova_evidence']),s(c['moldova_source']),s(c['jurisdictions_count']),s(c['jurisdictions_examples']),s(c['jurisdictions_source']),
   'Да' if pp.startswith('yes') else ('Сеть (неформально)' if 'network' in pp else 'Не найдена'),s(c['partner_program_details']),s(c['partner_source']),
   s(c['email']),s(c['phone']),s(c['contact_page']),s(c['address']),s(c['contact_source']),s(c['fit_notes']),
   next((v for k,v in RISK.items() if c['name'].startswith(k)),''),s(c.get('operating_evidence')),'Поисковая выдача по оф. домену; сайт не открыт']
  sc=c['scores'];vals+= [sc['client_relevance'],sc['partner_openness'],sc['moldova_opportunity'],sc['reachability'],sc['scale']]
  vals.append(f'=ROUND(Z{i}*Методика!$B$3+AA{i}*Методика!$B$4+AB{i}*Методика!$B$5+AC{i}*Методика!$B$6+AD{i}*Методика!$B$7,2)')
  for j,v in enumerate(vals,1):
    cell=r.cell(i,j,v);cell.font=BLUE if 26<=j<=30 else F;cell.alignment=WR
  r.cell(i,31).number_format='0.00'
  if c['website'].startswith('http'): r.cell(i,5).hyperlink=re.split(r'\s',c['website'])[0]
r.freeze_panes='D2';r.auto_filter.ref=f'A1:{L(len(cols))}{last}'
from openpyxl.formatting.rule import FormulaRule
r.conditional_formatting.add(f'A2:{L(len(cols))}{last}',FormulaRule(formula=['$A2<=5'],fill=TOP))
r.row_dimensions[1].height=40
# --- Группы
def group(title,flag,desc):
  g=wb.create_sheet(title);g['A1']=desc;g['A1'].font=B
  hdr=['Ранг','Компания','Сайт','Молдова','Доказательство / источник','Партнёрская программа','Email','Телефон','Итоговый балл']
  for j,h in enumerate(hdr,1):
    c=g.cell(2,j,h);c.font=H;c.fill=HF
  k=3
  for i,c in enumerate(rows,2):
    if (moldova(c)=='Да')!=flag: continue
    refs=[f"=Рейтинг!A{i}",f"=Рейтинг!C{i}",f"=Рейтинг!E{i}",f"=Рейтинг!H{i}",f'=Рейтинг!I{i}&" | "&Рейтинг!J{i}',f'=Рейтинг!N{i}&": "&Рейтинг!O{i}',f"=Рейтинг!Q{i}",f"=Рейтинг!R{i}",f"=Рейтинг!AE{i}"]
    for j,v in enumerate(refs,1):
      cc=g.cell(k,j,v);cc.font=GREEN;cc.alignment=WR
    g.cell(k,9).number_format='0.00';k+=1
  g['A%d'%(k+1)]=f'Всего: =COUNTA(B3:B{k-1})';g['A%d'%(k+1)]=f'=COUNTA(B3:B{k-1})';g['B%d'%(k+1)]='компаний в группе';g['A%d'%(k+1)].font=B
  for j,w in enumerate([6,28,26,12,60,60,26,22,9],1): g.column_dimensions[L(j)].width=w
  g.freeze_panes='C3'
group('Группа A – нет Молдовы',False,'Группа A: Молдова отсутствует или не найдена в опубликованном каталоге (зелёные ячейки берутся с листа «Рейтинг»)')
group('Группа B – есть Молдова',True,'Группа B: Молдова уже предлагается')
# --- TOP-5 letters
t=wb.create_sheet('TOP-5 письма')
t['A1']='TOP-5: черновики писем. СТАТУС: НЕ ОТПРАВЛЕНО, нужно утверждение. Полный текст в research/top5_letters.md';t['A1'].font=Font(name='Arial',size=11,bold=True,color='C00000')
hdr=['Ранг','Компания','Кому / канал','Тема письма','Ключевой аргумент','Что сверить до отправки','Статус']
for j,h in enumerate(hdr,1):
  c=t.cell(2,j,h);c.font=H;c.fill=HF
L5=[('Accace','Accace Circle, https://circle.accace.com/ (общий email не подтверждён)','Moldova legal partner for Accace clients | Legitimus, Chișinău','Молдова уже предлагается, рынки без собственных филиалов покрываются через партнёров (Accace Circle). Legitimus как юридический партнёр по Молдове или резервный','Есть ли у Accace собственный офис в Кишинёве'),
('BridgeWest','office@bridgewest.eu; Join Our Network на bridgewest.eu','Joining the BridgeWest network as your Moldova affiliate | Legitimus','Продаёт регистрацию в Молдове через независимых аффилиатов в 70+ странах; Legitimus как аффилиат по Молдове','Страница по Молдове и условия вступления в сеть'),
('Commitbiz','info@commitbiz.com; /contact-us','Moldova partner for Commitbiz | referral cooperation proposal','Есть офис в Бухаресте и формальный Referral Agreement; Молдовы нет. Взаимные рекомендации','Офис в Бухаресте; условия реферального соглашения'),
('Virtuzone','info@virtuzone.com; referral.virtuzone.com','Partnership proposal: Moldova legal support for Virtuzone / Ascentium clients','История работы с СНГ (Virtuzone CIS), Молдовы нет; клиенты из Молдовы в ОАЭ и местный юрист в Молдове для клиентов Virtuzone/Ascentium','Текущий статус Virtuzone CIS; контакт по партнёрствам'),
('Uniwide','Форма на /partner-programme/ или /contact/; +371 6611 8787','Uniwide Partner Programme | law firm in Moldova (referral and local agent)','Программа открыта для юрфирм (реферальная и субподрядная модели), работает через локальных агентов','Полный список юрисдикций (есть ли Молдова)')]
for k,(nm,to,subj,arg,chk) in enumerate(L5,3):
  i=next(ix for ix,c in enumerate(rows,2) if c['name'].startswith(nm))
  for j,v in enumerate([f'=Рейтинг!A{i}',f'=Рейтинг!C{i}',to,subj,arg,chk,'Черновик, не отправлено'],1):
    cc=t.cell(k,j,v);cc.font=GREEN if j<=2 else F;cc.alignment=WR
for j,w in enumerate([6,26,40,45,60,40,22],1): t.column_dimensions[L(j)].width=w
# --- Исключены
x=wb.create_sheet('Исключены')
for j,h in enumerate(['Компания','Причина исключения'],1):
  c=x.cell(1,j,h);c.font=H;c.fill=HF
ex=[(c['name'],c['drop_reason']) for c in d['dropped']]+[
('Offshore Company Corp (offshorecompany.com)','Не удалось ничего подтвердить по собственному сайту; заменена на Uniwide'),
('Startupr (startupr.com)','Не удалось подтвердить, что компания действует: в индексе только старые страницы (2012)'),
('Prime Corporate Services (primecorporateservices.com)','Домен принадлежит американской фирме из Юты (формирование компаний в США), а не кипрскому провайдеру'),
('Intercompany Solutions','Регистрирует компании только в Нидерландах'),
('Emirabiz; Flyingcolour','Регистрируют компании только в ОАЭ'),
('vzone.ae','Посторонний сайт (вейп-шоп), не Virtuzone; правильный домен virtuzone.com')]
for k,(a,b) in enumerate(ex,2):
  x.cell(k,1,a).font=F;x.cell(k,2,b).font=F;x.cell(k,2).alignment=WR
x.column_dimensions['A'].width=45;x.column_dimensions['B'].width=90
wb.move_sheet('Методика',offset=len(wb.sheetnames)-1)
wb.active=0
wb.save('/home/user/-/research/Legitimus_partners_research.xlsx')
