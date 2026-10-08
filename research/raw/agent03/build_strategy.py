import json,sys
sys.path.insert(0,'.')
from strategy import S,MODELS,DECISIONS
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter as L
A=[c for c in json.load(open('/home/user/-/research/raw/agent02/final.json')) if c['status'].startswith('TOP-10')]
def key(c): return next(k for k in S if c['name'].startswith(k))
F=Font(name='Arial',size=10);B=Font(name='Arial',size=10,bold=True);H=Font(name='Arial',size=10,bold=True,color='FFFFFF')
HF=PatternFill('solid',fgColor='1F3864');WR=Alignment(wrap_text=True,vertical='top');Y=PatternFill('solid',fgColor='FFFF00')
def j(v): return '; '.join(map(str,v)) if isinstance(v,list) else ('' if v is None else str(v))
def hdr(ws,cols,row=1):
  for i,(h,w) in enumerate(cols,1):
    c=ws.cell(row,i,h);c.font=H;c.fill=HF;c.alignment=Alignment(wrap_text=True,vertical='center');ws.column_dimensions[L(i)].width=w
  ws.row_dimensions[row].height=36
def put(ws,r,vals):
  for i,v in enumerate(vals,1):
    c=ws.cell(r,i,v);c.font=F;c.alignment=WR
wb=Workbook()
# README
s=wb.active;s.title='Сводка'
lines=[('LEGITIMUS · TOP-10 Partnership Strategy (Агент 03: International Partnership Strategist)',True),
('Основа: проверенные данные агента 02 (Verified Partners Database.xlsx, final.json), дата проверки 2026-10-08. Новое исследование не проводилось; поиск только для проверки существенных коммерческих предположений (лист «Допущения»).',False),
('Уровни достоверности: High = официальный сайт или реестр; Medium = официальный домен или пресса; Low = сторонний источник. ASSUMPTION = допущение. [TO CONFIRM] = требует подтверждения LEGITIMUS.',False),
('Письма не отправлялись, переговоры не велись. Листы «Outreach Briefs» предназначены для следующего агента (Outreach Agent) и используются только после утверждения решений на листе «Решения MP».',False),('',False),
('Распределение моделей:',True)]
for i,(t,b) in enumerate(lines,1):
  s.cell(i,1,t).font=B if b else F;s.cell(i,1).alignment=WR;s.merge_cells(start_row=i,start_column=1,end_row=i,end_column=6);s.row_dimensions[i].height=30 if not b else 18
for k,(m,name) in enumerate([('A','Model A — New Jurisdiction Partner'),('B','Model B — Local Execution Partner'),('C','Model C — Long-Term Legal Desk')],7):
  s.cell(k,1,name).font=F;s.cell(k,2,f"=COUNTIF('TOP-10 стратегия'!$E:$E,\"{m}\")").font=F
s.column_dimensions['A'].width=60;s.column_dimensions['B'].width=10
# Models
m=wb.create_sheet('Модели A-B-C')
hdr(m,[('Параметр',26),('Model A — New Jurisdiction Partner',55),('Model B — Local Execution Partner',55),('Model C — Long-Term Legal Desk',55)])
rows=[('Для кого','for_whom'),('Предложение','offer'),('Выгоды партнёра','benefits'),('Процесс передачи заявки','handover'),('Распределение обязанностей','roles'),('Коммерческие вопросы для согласования','commercial')]
for r,(lab,k) in enumerate(rows,2):
  vals=[lab]
  for mm in 'ABC':
    v=MODELS[mm][k]
    if k=='roles': v='\n'.join(f'{a}: {b}' for a,b in v)
    elif isinstance(v,list): v='\n'.join(('• '+x) if not x[:2].strip('.').isdigit() else x for x in v)
    vals.append(v)
  put(m,r,vals);m.cell(r,1).font=B
# Strategy
t=wb.create_sheet('TOP-10 стратегия')
cols=[('Ранг (агент 02)',7),('Компания',22),('Приоритет',14),('Балл /100',8),('Модель',7),('1. Профиль компании',45),('2. Почему ей может быть интересна Молдова',45),('3. Модель и обоснование',30),('4. Какую проблему решает LEGITIMUS',40),('5. Услуги в первом предложении',45),('6. Адресат',40),('7. Главный аргумент для первого контакта',45),('8. Возражения и ответы',60),('9. Следующее действие',45)]
hdr(t,cols)
for r,c in enumerate(A,2):
  d=S[key(c)];mm=d['model']
  put(t,r,[r-1,c['name'],c['status'].replace('TOP-10 · ',''),c['total'],mm,d['profile'],d['why_md'],f"{MODELS[mm]['name']}. {MODELS[mm]['for_whom']}",d['problem'],d['services'],d['addressee'],d['argument'],
   '\n'.join(f'«{a}» → {b}' for a,b in d['objections']),d['next']])
t.freeze_panes='C2';t.auto_filter.ref=f'A1:{L(len(cols))}{len(A)+1}'
# Objections
o=wb.create_sheet('Возражения и ответы')
hdr(o,[('Компания',24),('Модель',7),('Возможное возражение',50),('Рекомендуемый ответ',70)])
r=2
for c in A:
  d=S[key(c)]
  for a,b in d['objections']: put(o,r,[c['name'],d['model'],a,b]);r+=1
# Outreach briefs
br=wb.create_sheet('Outreach Briefs')
cols=[('Компания',22),('Сайт',24),('Модель',7),('Адресат (имя, роль, достоверность)',40),('Публичный канал',38),('Проверенные факты (EN, с достоверностью)',70),('Конкретное предложение',50),('Главный аргумент',45),('Рекомендуемый тон',30),('Не говорить / не утверждать',40),('Сверить перед письмом',40),('Источники',60),('Статус',16)]
hdr(br,cols)
for r,c in enumerate(A,2):
  d=S[key(c)];ct=c['contacts']
  facts=[f"Moldova: {c['moldova']['value']} (delivery: {c['moldova']['delivery']}) [{c['moldova']['confidence']}] — {j(c['moldova']['evidence'])}",
         f"Multi-country: {j(c['multi_country'].get('evidence'))} [{c['multi_country'].get('confidence')}]",
         f"Law-firm cooperation: {j(c['law_firm_cooperation']['value'])} — {j(c['law_firm_cooperation']['evidence'])} [{c['law_firm_cooperation']['confidence']}]"]
  dm='\n'.join(f"{x['name']} — {x['title']} ({x['confidence']})" for x in c['decision_makers']) or 'not found'
  src=[]
  for k in ['exists','multi_country','moldova','law_firm_cooperation','long_term_services','contacts']:
    v=c.get(k,{}).get('sources') if isinstance(c.get(k),dict) else None
    if v: src+= v if isinstance(v,list) else [v]
  src+=[x['source'] for x in c['decision_makers'] if x.get('source')]
  src=list(dict.fromkeys(s for s in src if isinstance(s,str) and s.strip()))
  checks=[x for x in c['audit'] if x['status'] in ('needs_check','unconfirmed','outdated')]
  put(br,r,[c['name'],c['website'],d['model'],dm+'\n\nРекомендовано: '+d['addressee'],'\n'.join(filter(None,[j(ct.get('email')),j(ct.get('phone')),j(ct.get('contact_page'))]))+f"\n(current: {ct.get('current')}, confidence: {ct.get('confidence')})",
   '\n'.join(facts),d['services'],d['argument'],d['tone'],d['dont']+' Не указывать цены и не делать предложений по вознаграждению.',
   '\n'.join(f"• {x['claim']}: {x['note']}" for x in checks[:6]) or '—','\n'.join(src),'Ждёт утверждения MP'])
br.freeze_panes='B2'
# Assumptions
a=wb.create_sheet('Допущения')
hdr(a,[('Допущение / вопрос',45),('Результат проверки',60),('Статус',18),('Источник',60),('Влияние на стратегию',45)])
AS=[('Кодекс этики адвокатов РМ разрешает или запрещает раздел гонорара и реферальные вознаграждения','В открытых источниках найдены только общие нормы: гонорар устанавливается свободно по соглашению с клиентом, должен быть справедливым, нарушение Кодекса является дисциплинарным проступком. Специальная норма о разделе гонорара или комиссиях не найдена.','НЕ ПОДТВЕРЖДЕНО','https://crjm.org/wp-content/uploads/2002/12/Codul-deontologic-avocati.pdf','Реферальные вознаграждения и white-label НЕ предлагать до заключения (Решение 1)'),
('MyAccounting: партнёр по бухгалтерии в Кишинёве','Публичных сведений о компании MyAccounting в Молдове не найдено (поиск 2026-10-08).','НЕ ПОДТВЕРЖДЕНО','Поиск без результатов','Реквизиты, юр. форма и договор с LEGITIMUS: [TO CONFIRM] (Решение 3)'),
('Закон 195/2024 о защите персональных данных действует','Вступил в силу 23.08.2026, заменил Закон 133/2011; надзор: Национальный центр по защите персональных данных. Источники расходятся в деталях санкций.','ПОДТВЕРЖДЕНО (Medium)','https://moldpres.md/eng/society/new-rules-on-personal-data-protection-enter-into-force-in-moldova-as-of-august-23','Использовать как аргумент для Data Protection; без цифр штрафов'),
('Переговоры РМ о вступлении в ЕС продвигаются','Открыты кластер 1 «Fundamentals» (15.06.2026) и кластер 6 «External Relations» (14.07.2026); остальные пока не открыты.','ПОДТВЕРЖДЕНО (High)','https://www.consilium.europa.eu/en/policies/moldova','Аргумент «почему Молдова сейчас»; без прогнозов сроков вступления'),
('Провайдеры без офиса в Кишинёве используют местного исполнителя','Логический вывод из данных агента 02 (Accace, BridgeWest, Healy); не подтверждён ими напрямую.','ASSUMPTION','research/raw/agent02/final.json','Формулировать как вопрос в первом контакте'),
('Клиенты провайдеров из Румынии и Украины интересуются Молдовой','Не подтверждено данными.','ASSUMPTION','—','Задавать вопрос, не утверждать'),
('Регистрация в ASP требует участия адвоката','Не подтверждено; не использовать как аргумент.','НЕ ПОДТВЕРЖДЕНО','—','Аргумент строить на качестве, ответственности и адвокатской тайне')]
for r,row in enumerate(AS,2): put(a,r,list(row))
# Decisions
dsh=wb.create_sheet('Решения MP')
hdr(dsh,[('#',4),('Решение',30),('Суть и почему важно',70),('Когда нужно',30),('Что утвердить',50),('Решение MP',30)])
for r,(n,why,when,what) in enumerate(DECISIONS,2):
  put(dsh,r,[r-1,n,why,when,what,'']);dsh.cell(r,6).fill=Y
wb.save('/home/user/-/research/agent03/TOP10_Partnership_Strategy.xlsx')
