import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule
A=json.load(open('final.json'))
F=Font(name='Arial',size=10);B=Font(name='Arial',size=10,bold=True);H=Font(name='Arial',size=10,bold=True,color='FFFFFF')
BLUE=Font(name='Arial',size=10,color='0000FF');GREEN=Font(name='Arial',size=10,color='008000')
HF=PatternFill('solid',fgColor='1F3864');WR=Alignment(wrap_text=True,vertical='top')
P1=PatternFill('solid',fgColor='C6EFCE');P2=PatternFill('solid',fgColor='FFF2CC');HOLD=PatternFill('solid',fgColor='F8CBAD');EX=PatternFill('solid',fgColor='D9D9D9')
def j(v):
  if isinstance(v,list): return '\n'.join(map(str,v))
  return '' if v is None else str(v)
def hdr(ws,cols,row=1):
  for i,(h,w) in enumerate(cols,1):
    c=ws.cell(row,i,h);c.font=H;c.fill=HF;c.alignment=Alignment(wrap_text=True,vertical='center');ws.column_dimensions[L(i)].width=w
  ws.row_dimensions[row].height=42
K=['profile_fit','recurring_client_potential','local_legal_partner_cooperation','moldova_commercial_interest','decision_maker_accessibility']
wb=Workbook()
# ---------- Methodology / summary
s=wb.active;s.title='Сводка'
s['A1']='LEGITIMUS · Verified Partners Database (Агент 02: Partnership Intelligence Analyst)';s['A1'].font=Font(name='Arial',size=13,bold=True)
s['A2']='Дата проверки: 2026-10-08. Источник для проверки: база агента 01 (Legitimus_partners_research.xlsx).';s['A2'].font=F
s['A4']='Статус';s['B4']='Кол-во';s['A4'].font=B;s['B4'].font=B
sts=['TOP-10 · Приоритет 1','TOP-10 · Приоритет 2','Отложено: требуется комплаенс-проверка','Резерв','Исключено']
for i,t in enumerate(sts,5):
  s[f'A{i}']=t;s[f'B{i}']=f'=COUNTIF(\'Verified Database\'!$E:$E,A{i})';s[f'A{i}'].font=F;s[f'B{i}'].font=F
s['A10']='Всего';s['B10']='=SUM(B5:B9)';s['A10'].font=B;s['B10'].font=B
s['A12']='Шкала оценки (100 баллов)';s['A12'].font=B
crit=[('Соответствие целевому профилю LEGITIMUS',20),('Потенциал клиентов на постоянное обслуживание',30),('Возможность сотрудничества с местным юр. партнёром',20),('Коммерческий интерес к Молдове',15),('Доступность лиц, принимающих решения',15)]
for i,(n,m) in enumerate(crit,13): s[f'A{i}']=n;s[f'B{i}']=m;s[f'A{i}'].font=F;s[f'B{i}'].font=F
s['A18']='Максимум';s['B18']='=SUM(B13:B17)';s['A18'].font=B;s['B18'].font=B
notes=['Уровни достоверности: High = текст официального сайта или реестра; Medium = сниппет с официального домена или из авторитетной прессы; Low = сторонний каталог или косвенные данные.',
'«ASSUMPTION» = аналитическое допущение, не факт. «not found» означает, что сведения не найдены, а не что их нет.',
'ОГРАНИЧЕНИЕ: сетевая политика среды блокировала открытие сайтов компаний. Проверка шла через поисковую выдачу (site:-запросы к официальным доменам, пресс-релизы, отраслевые СМИ). Перед контактом сверьте ключевые факты на сайтах.',
'Баллы (синие) — экспертная оценка агента 02 с обоснованием в соседних столбцах. Их можно менять: итог и ранг пересчитаются.',
'GSL набрал 65 баллов, но выведен из TOP-10 до формальной проверки по санкционным спискам (решение аналитика). Его место занял Virtuzone (59).',
'Никаких контактов с компаниями не было: письма не отправлялись, формы не заполнялись.']
for i,t in enumerate(notes,20):
  s[f'A{i}']=t;s[f'A{i}'].font=F;s[f'A{i}'].alignment=WR;s.merge_cells(f'A{i}:H{i}');s.row_dimensions[i].height=32
s.column_dimensions['A'].width=55;s.column_dimensions['B'].width=10
# ---------- Database
d=wb.create_sheet('Verified Database')
cols=[('Ранг',6),('Компания',26),('Сайт',24),('Штаб-квартира',24),('Статус',22),
('Профиль /20',8),('Пост. клиенты /30',8),('Локальный партнёр /20',9),('Интерес к Молдове /15',9),('Доступность ЛПР /15',9),('ИТОГО /100',8),('Оценка агента 01 (×20, индикативно)',10),('Δ к агенту 01',8),
('Обоснование оценок',60),('Существует / активна',35),('Специализация',35),('Регистрация в нескольких странах',35),
('Молдова',12),('Как оказывается в Молдове',14),('Доказательства по Молдове',50),('Достоверность (Молдова)',10),('Источники (Молдова)',40),
('Сотрудничество с юрфирмами',45),('Долгосрочные услуги',40),('Признаки разовых продаж',40),
('Email',28),('Телефон',28),('Страница контактов',30),('Контакты актуальны?',10),('Достоверность (контакты)',10),
('ЛПР (имя, должность, источник, достоверность)',55),('Причина исключения',40),('Дата проверки',11)]
hdr(d,cols)
n=len(A);last=n+1
for r,c in enumerate(A,2):
  sc=c['scores']
  why='\n'.join(f"{k}: {sc[k]['why']}" for k in K)
  ex=lambda o:f"{j(o.get('value'))} — {j(o.get('evidence',''))} [{o.get('confidence','')}] {j(o.get('sources',''))}"
  dm='\n'.join(f"{x['name']} — {x['title']} ({x.get('confidence','')}) {x.get('source','')}" for x in c['decision_makers']) or 'not found'
  ct=c['contacts']
  vals=[f'=RANK(K{r},$K$2:$K${last})+COUNTIF($K$2:K{r},K{r})-1',c['name'],c['website'],c['hq'],c['status']]+[int(sc[k]['score']) for k in K]+\
   [f'=SUM(F{r}:J{r})',c['a1_score100'],f'=K{r}-L{r}',why,ex(c['exists']),j(c['specialization']),ex(c['multi_country']),
    c['moldova']['value'],c['moldova'].get('delivery'),j(c['moldova']['evidence']),c['moldova']['confidence'],j(c['moldova']['sources']),
    ex(c['law_firm_cooperation']),f"{j(c['long_term_services'].get('value'))} — {j(c['long_term_services'].get('evidence'))} {j(c['long_term_services'].get('sources'))}",j(c['one_off_signals']),
    j(ct.get('email')),j(ct.get('phone')),j(ct.get('contact_page')),j(ct.get('current')),j(ct.get('confidence')),dm,
    c['exclude']['reason'] if c['exclude']['value'] else '',c['check_date']]
  for i,v in enumerate(vals,1):
    cell=d.cell(r,i,v);cell.font=BLUE if 6<=i<=10 else F;cell.alignment=WR
  if c['website'].startswith('http'): d.cell(r,3).hyperlink=c['website'].split()[0]
rng=f'A2:{L(len(cols))}{last}'
for f,fill in [('LEFT($E2,13)="TOP-10 · Прио"',None)]: pass
d.conditional_formatting.add(rng,FormulaRule(formula=['$E2="TOP-10 · Приоритет 1"'],fill=P1))
d.conditional_formatting.add(rng,FormulaRule(formula=['$E2="TOP-10 · Приоритет 2"'],fill=P2))
d.conditional_formatting.add(rng,FormulaRule(formula=['LEFT($E2,8)="Отложено"'],fill=HOLD))
d.conditional_formatting.add(rng,FormulaRule(formula=['$E2="Исключено"'],fill=EX))
d.freeze_panes='C2';d.auto_filter.ref=f'A1:{L(len(cols))}{last}'
# ---------- TOP-10
t=wb.create_sheet('TOP-10')
tc=[('Ранг в базе',7),('Компания',26),('Статус',20),('ИТОГО /100',8),('Молдова',12),('Формат сотрудничества',30),('Что предложить',50),('Кому обращаться (ЛПР)',45),('Публичный канал',40),('Возможные возражения',50),('Первый шаг',50)]
hdr(t,tc)
top=[(r,c) for r,c in enumerate(A,2) if c['status'].startswith('TOP-10')]
for k,(r,c) in enumerate(top,2):
  dd=c['dossier'];ct=c['contacts']
  vals=[f"='Verified Database'!A{r}",f"='Verified Database'!B{r}",f"='Verified Database'!E{r}",f"='Verified Database'!K{r}",f"='Verified Database'!R{r}",
   j(dd['preferred_format']),j(dd['what_to_offer']),'\n'.join(f"{x['name']} — {x['title']} ({x.get('confidence','')})" for x in c['decision_makers']) or 'not found',
   '\n'.join(filter(None,[j(ct.get('email')),j(ct.get('contact_page'))])),j(dd['objections']),j(dd['first_step'])]
  for i,v in enumerate(vals,1):
    cell=t.cell(k,i,v);cell.font=GREEN if i<=5 else F;cell.alignment=WR
t.freeze_panes='C2'
# ---------- Audit
a=wb.create_sheet('Аудит агента 01')
a['A1']='Сводка по статусам утверждений агента 01';a['A1'].font=B
labels=[('confirmed','Подтверждено'),('erroneous','Ошибочно'),('unconfirmed','Не подтверждено'),('outdated','Устарело (в т.ч. контакты)'),('duplicate','Дубликат'),('needs_check','Требует доп. проверки')]
rows=[]
for c in A:
  for x in c['audit']: rows.append((c['name'],x.get('claim'),x.get('status'),x.get('note'),j(x.get('source'))))
start=12;endr=start+len(rows)
for i,(k,lab) in enumerate(labels,2):
  a[f'A{i}']=lab;a[f'B{i}']=k;a[f'C{i}']=f'=COUNTIF($C${start+1}:$C${endr},B{i})'
  for col in 'ABC': a[f'{col}{i}'].font=F
a['A8']='Всего утверждений';a['C8']=f'=SUM(C2:C7)';a['A8'].font=B;a['C8'].font=B
a['A9']='Компании, исключённые агентом 02: см. лист «Verified Database», статус «Исключено». Требуют доп. проверки: строки со статусом needs_check / unconfirmed.';a['A9'].font=F
for i,(h,w) in enumerate([('Компания',28),('Утверждение агента 01',50),('Статус',14),('Комментарий агента 02',70),('Источник',50)],1):
  c=a.cell(start,i,h);c.font=H;c.fill=HF;a.column_dimensions[L(i)].width=w
for k,row in enumerate(rows,start+1):
  for i,v in enumerate(row,1):
    cell=a.cell(k,i,j(v));cell.font=F;cell.alignment=WR
a.auto_filter.ref=f'A{start}:E{endr}';a.freeze_panes=f'A{start+1}'
# ---------- Excluded summary
e=wb.create_sheet('Исключены и отложены')
hdr(e,[('Компания',30),('Статус',24),('ИТОГО /100',9),('Причина',90)])
k=2
for r,c in enumerate(A,2):
  if c['status'].startswith('TOP-10'): continue
  reason=c['exclude']['reason'] if c['exclude']['value'] else ('Санкционный и репутационный риск: штаб-квартира в Москве, российская клиентура. До контакта нужна формальная проверка по санкционным спискам и одобрение партнёра.' if c['name'].startswith('GSL') else 'Резерв: невысокий балл, в основном разовые регистрации; вернуться после TOP-10')
  for i,v in enumerate([f"='Verified Database'!B{r}",f"='Verified Database'!E{r}",f"='Verified Database'!K{r}",reason],1):
    cell=e.cell(k,i,v);cell.font=GREEN if i<=3 else F;cell.alignment=WR
  k+=1
wb.save('/home/user/-/research/agent02/Verified Partners Database.xlsx')
